"""Gestaltete, ausfüllbare Word-Unterlagen aus den fachlichen Markdown-Quellen.

Mit gebündeltem Python ausführen und anschließend beide DOCX vollständig rendern.
Gelb hinterlegte Textfelder sind editierbar; Quellenmarker bleiben unverändert.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ORDNER = Path(__file__).resolve().parents[1]
VERSION = "1.0"
BLAU = "24485B"
GELB = "FFF2CC"

# Die Übersicht enthält Entscheidungen und Nachweise, keine vorweggenommenen Antworten.
OFFENE_PUNKTE = [
    ("Regierung", "ERGÄNZEN", "Sachgebiet 40.1 und Doris Sippel stehen fest. Ergänzen: Direktkontakt, Vertretung, Schulämter, befugte Entscheidungsstelle und Aktenzeichen.", "1 und 7", "1 und 7"),
    ("Regierung", "BESTÄTIGEN", "Beschriebene Aufgaben, Auswertungen und Rechte bestätigen; tatsächliche FIBS-Quelle und Nutzung klären. Datenfelder sind bereits beschrieben.", "TOM 02 und 05", "VVT 01 und 02"),
    ("Regierung und Registratur", "FESTLEGEN", "Archivfrist und Fristbeginn, öffentliche Anzeigedauer sowie Fristen für Konten, Nachweise und Exporte. Die operative 400-Tage-Regel steht.", "TOM 06 und 07", "VVT 04 und 6"),
    ("Vertragsstelle und IT", "BELEGEN", "AVV und Anlagen ablegen; konkrete Serverinstanz Nürnberg, Portaladresse und gegebenenfalls weitere Dienstleister eintragen.", "2 und TOM 04", "1 und VVT 03"),
    ("Projektteam und Administration", "PRÜFEN", "Zugänge, BdB-MFA, Bezirksrechte und Löschlauf auf dem Server prüfen. Vollbackup: age-Schlüssel, UNC-Pfad, Rechte, Vertretung, bestätigte Ablage und Wiederherstellung belegen; technische Logs festlegen.", "TOM 01 bis 09", "VVT 03 und 6"),
    ("Regierung, IT und Datenschutz", "ABSCHLIESSEN", "Risikovorprüfung, Meldewege und Zuständigkeit für Auskünfte festhalten; Datenschutzinformationen und Betriebsfreigabe vervollständigen.", "TOM 09 und 10", "7"),
]


def hinterlegen(lauf, farbe=GELB):
    eig = lauf._r.get_or_add_rPr()
    feld = OxmlElement("w:shd")
    feld.set(qn("w:fill"), farbe)
    eig.append(feld)


def feldtext(absatz, text, fett=False):
    """Nur echte Platzhalter hervorheben, keine Quellen oder Musterkennzeichnung."""
    for teil in re.split(r"(\[[^\]]+\])", text):
        if not teil:
            continue
        lauf = absatz.add_run(teil)
        lauf.bold = True if fett else None
        if teil.startswith("[") and not re.fullmatch(r"\[Q\d+(?:,\s*Q\d+)*\]", teil) and teil != "[Schulamtsbezirk]":
            lauf.text = f"[AUSFÜLLEN: {teil[1:-1]}]"
            lauf.bold = True
            hinterlegen(lauf)


def text_einfuegen(absatz, text):
    muster = r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)"
    for teil in re.split(muster, text):
        if not teil:
            continue
        if teil.startswith("**"):
            feldtext(absatz, teil[2:-2], fett=True)
        elif teil.startswith("`"):
            feldtext(absatz, teil[1:-1])
        elif re.match(r"\[[^\]]+\]\(", teil):
            treffer = re.fullmatch(r"\[([^\]]+)\]\((.+)\)", teil)
            if not treffer[2].startswith("http"):
                feldtext(absatz, f"{treffer[1]} ({treffer[2]})")
                continue
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), absatz.part.relate_to(treffer[2], RT.HYPERLINK, is_external=True))
            lauf = OxmlElement("w:r")
            eig = OxmlElement("w:rPr")
            farbe = OxmlElement("w:color"); farbe.set(qn("w:val"), BLAU); eig.append(farbe)
            lauf.append(eig)
            t = OxmlElement("w:t"); t.text = treffer[1]; lauf.append(t)
            link.append(lauf); absatz._p.append(link)
        else:
            feldtext(absatz, teil)


def tabelle(dokument, inhalt, breiten, schrift=9.5):
    tab = dokument.add_table(rows=0, cols=len(breiten))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for spalte, breite in zip(tab.columns, breiten):
        spalte.width = Cm(breite)
    grenzen = OxmlElement("w:tblBorders")
    for seite in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        g = OxmlElement(f"w:{seite}"); g.set(qn("w:val"), "single"); g.set(qn("w:sz"), "4"); g.set(qn("w:color"), "D9D9D9"); grenzen.append(g)
    tab._tbl.tblPr.append(grenzen)
    for nummer, inhalte in enumerate(inhalt):
        zeile = tab.add_row()
        eig = zeile._tr.get_or_add_trPr()
        eig.append(OxmlElement("w:cantSplit"))
        if nummer == 0:
            eig.append(OxmlElement("w:tblHeader"))
        for zelle, wert, breite in zip(zeile.cells, inhalte, breiten):
            zelle.width = Cm(breite)
            zelle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            zelleneig = zelle._tc.get_or_add_tcPr()
            innen = OxmlElement("w:tcMar")
            for seite, abstand in [("top", 100), ("left", 115), ("bottom", 100), ("right", 115)]:
                m = OxmlElement(f"w:{seite}"); m.set(qn("w:w"), str(abstand)); m.set(qn("w:type"), "dxa"); innen.append(m)
            zelleneig.append(innen)
            fill = BLAU if nummer == 0 else ("F2F6F8" if nummer % 2 == 0 else "FFFFFF")
            schattierung = OxmlElement("w:shd"); schattierung.set(qn("w:fill"), fill); zelleneig.append(schattierung)
            p = zelle.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = nummer == 0
            text_einfuegen(p, wert)
            for lauf in p.runs:
                lauf.font.size = Pt(schrift)
                if nummer == 0:
                    lauf.bold = True; lauf.font.color.rgb = RGBColor(255, 255, 255)
    p = dokument.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(2)
    return tab


def tabellenzeilen(dokument, zeilen):
    inhalt = [[z.strip() for z in x.strip().strip("|").split("|")] for x in zeilen]
    inhalt = [x for x in inhalt if not all(re.fullmatch(r"[-: ]+", y) for y in x)]
    if len(inhalt[0]) == 2:
        breiten = [3.8, 13.2]
    elif inhalt[0][0] == "Datenkategorie":
        breiten = [4.1, 8.1, 4.8]
    elif inhalt[0][1] == "Federführung":
        breiten = [7.3, 4.2, 5.5]
    else:
        breiten = [4.3, 5.8, 6.9]
    tabelle(dokument, inhalt, breiten)


def startseite(dokument, kurz):
    titel = "Technische und\norganisatorische Maßnahmen" if kurz == "TOM" else "Verzeichnis der\nVerarbeitungstätigkeiten"
    dokument.add_paragraph(titel, style="Title")
    dokument.add_paragraph("Fortbildungsportal Schwaben", style="Subtitle")
    meta = "TOM 01" if kurz == "TOM" else "VVT 01 bis 04"
    dokument.add_paragraph(f"{meta}  ·  Version {VERSION}  ·  24. September 2026", style="Metadaten")
    p = dokument.add_paragraph()
    p.add_run("Zweck des Portals. ").bold = True
    p.add_run("Die Regierung von Schwaben und die Beratung digitale Bildung planen und bewerben dienstliche Fortbildungen. Sie verwalten die zugehörigen Referenten und werten das Angebot nach Schuljahr und Schulamtsbezirk aus. Das Portal verarbeitet dafür dienstliche Angaben; personenbezogene Teilnehmerlisten sind nicht vorgesehen.")
    p = dokument.add_paragraph()
    p.add_run("Wozu dieses Dokument dient. ").bold = True
    p.add_run("Die TOM erklären, wie Personen und ihre Daten geschützt werden und wer diese Maßnahmen im Betrieb verantwortet." if kurz == "TOM" else "Das VVT erklärt, wofür welche personenbezogenen Daten benötigt werden, wer sie sehen darf und wie lange sie aufbewahrt werden. Die zugehörigen TOM beschreiben den Schutz dieser Daten.")
    dokument.add_paragraph("Bereits geschaffene Schutzmaßnahmen", style="Heading 1")
    tabelle(dokument, [
        ["Schutz für die Betroffenen", "Was bereits erreicht ist"],
        ["Auswertungen bleiben im zuständigen Bezirk", "**Im Tool umgesetzt.** BdBs sehen ihre zugewiesenen Bezirke; nur ausdrücklich berechtigte Regierungskonten haben fachlichen Gesamtzugriff. Auch Downloads prüfen die Rechte."],
        ["Öffentliche Namen bleiben freiwillig", "**Im Tool umgesetzt.** Ohne eigene Zustimmung keine Namensanzeige im Referentenfeld. Zustimmung und Widerruf werden dokumentiert und bei neuen Abrufen berücksichtigt."],
        ["Verwaltungskonten sind zusätzlich geschützt", "**Im Tool umgesetzt.** Für Regierungs-, BdB-/Verwaltungs- und Redaktionskonten ist neben dem Passwort ein zweiter Anmeldefaktor erforderlich."],
        ["Schuljahre werden getrennt behandelt", "**Im Tool umgesetzt.** Nach 400 Tagen ab Schuljahresende endet der reguläre Zugriff. Ein begrenztes Archivpaket und die bestätigte Übergabe bereiten die Löschung vor."],
        ["Weniger Weitergabe beim Seitenaufruf", "**Im Tool umgesetzt.** Keine eingebundenen fremden Analysewerkzeuge, Schriftarten oder Inhalte; öffentliche Referentenfelder sind auf freigegebene Namen begrenzt."],
        ["Betrieb und Ablagen sind bestimmt", "**Für den Betrieb festgelegt.** Server bei netcup in Nürnberg; tägliche verschlüsselte Vollsicherung wird durch RvS mit MFA auf das Regierungslaufwerk übernommen; gesondertes Archiv bei Regierung beziehungsweise Schulamt. HTTPS ist laut Betreiber vorhanden."],
    ], [5.3, 11.7], schrift=9.5)
    p = dokument.add_paragraph(style="Lesehilfe")
    p.add_run("So ist der Stand zu verstehen: ").bold = True
    p.add_run("Im Tool umgesetzt bedeutet: Die Funktion ist in der vorliegenden Anwendung vorhanden. Ihre Einrichtung und Wirksamkeit auf dem vorgesehenen Server werden vor Betriebsfreigabe bestätigt. Festgelegte Ablagen sind noch kein Nachweis erfolgreicher Sicherungen.")
    p = dokument.add_paragraph(style="Lesehilfe")
    p.add_run("Begrenzte Planungsausnahme: ").bold = True
    p.add_run("Zur Terminabstimmung dürfen planungsberechtigte Personen Veranstaltungsdaten anderer Bezirke sehen. Fremde Auswertungen, Nachbereitungsnotizen und Referentenkontakte gehören nicht zu dieser Ausnahme. Einzelheiten stehen in TOM 02 beziehungsweise VVT 01.")
    p = dokument.add_paragraph("Wer die restlichen Angaben ergänzt", style="Heading 1")
    p.paragraph_format.page_break_before = True
    p = dokument.add_paragraph()
    p.add_run("Die technische Grundlage ist vorhanden. ").bold = True
    p.add_run("Bereits beantwortet sind Datenumfang, Rechte, freiwillige Namensanzeige, BdB-MFA und 400-Tage-Regel. Sachgebiet 40.1 ist eingetragen. Das Projektteam ergänzt auch die technischen Betriebsnachweise. Die ausführenden Personen werden benannt; vorhandene behördliche Regelungen und Prüfprotokolle können genutzt werden.")
    zeilen = [["Wer", "Noch zu tun", "Fundstelle"]]
    for bereich, art, text, tom, vvt in OFFENE_PUNKTE:
        zeilen.append([bereich, f"**{art}**\n{text}", tom if kurz == "TOM" else vvt])
    tabelle(dokument, zeilen, [3.4, 10.2, 3.4], schrift=9)
    p = dokument.add_paragraph(style="Lesehilfe")
    p.add_run("So bearbeiten Sie die Vorlage: ").bold = True
    p.add_run("Gelb hinterlegte ")
    f = p.add_run("[AUSFÜLLEN: …]"); f.bold = True; hinterlegen(f)
    p.add_run(" durch die verbindliche Angabe ersetzen. Entscheidungen und Nachweise in den genannten Abschnitten dokumentieren. Anschließend Markierungen entfernen und die Freigabe auf der letzten Fachseite ausfüllen.")
    p = dokument.add_paragraph(style="Lesehilfe")
    p.add_run("Für die Freigabe wesentlich: ").bold = True
    p.add_run("Es geht um klare Verantwortlichkeiten, begründete Fristen und einen geprüften Betriebsstart. Die Datenschutzbeauftragte berät; die verantwortliche Stelle entscheidet über den Betrieb. Ein kurzer Prüfvermerk mit Datum, Ergebnis und zuständiger Person kann mehrere Punkte gemeinsam dokumentieren.")


def erzeugen(stamm):
    dokument = Document()
    s = dokument.sections[0]
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.top_margin = Cm(1.9); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2); s.right_margin = Cm(2)
    s.header_distance = Cm(.8); s.footer_distance = Cm(.8)
    for rand in dokument.styles.element.xpath(".//w:pBdr"):
        rand.getparent().remove(rand)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        stil = dokument.styles[name]
        stil.font.name = "Arial"; stil.font.color.rgb = RGBColor(0, 0, 0)
        stil.font.size = Pt(10.5)
        stil.paragraph_format.space_after = Pt(6)
        stil.paragraph_format.line_spacing = 1.07
        stil.paragraph_format.widow_control = True
    for name, groesse in [("Title", 25), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 11.5)]:
        stil = dokument.styles[name]; stil.font.size = Pt(groesse); stil.font.bold = True
        stil.paragraph_format.space_before = Pt(12 if name == "Heading 2" else 0)
        stil.paragraph_format.space_after = Pt(8)
        stil.paragraph_format.keep_with_next = True
    dokument.styles["Subtitle"].font.italic = False
    dokument.styles["Title"].paragraph_format.line_spacing = 1.06
    dokument.styles["List Bullet"].paragraph_format.space_after = Pt(5)
    for name, groesse in [("Metadaten", 9), ("Lesehilfe", 9), ("Quellen", 9)]:
        stil = dokument.styles.add_style(name, 1)
        stil.base_style = dokument.styles["Normal"]
        stil.font.name = "Arial"; stil.font.size = Pt(groesse)
        stil.paragraph_format.line_spacing = 1.08
        stil.paragraph_format.space_after = Pt(6)
    kurz = "TOM" if stamm.startswith("TOM") else "VVT"
    kopf = s.header.paragraphs[0]
    kopf.paragraph_format.tab_stops.clear_all()
    kopf.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_ALIGN_PARAGRAPH.RIGHT)
    kopf.text = f"FORTBILDUNGSPORTAL SCHWABEN\t{kurz}  ·  {VERSION}"
    kopf.runs[0].font.name = "Arial"; kopf.runs[0].font.size = Pt(8); kopf.runs[0].font.color.rgb = RGBColor(0,0,0)
    fuss = s.footer.paragraphs[0]
    fuss.paragraph_format.tab_stops.clear_all()
    fuss.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_ALIGN_PARAGRAPH.RIGHT)
    fuss.add_run("Freigabe offen  ·  Ausfüllfelder gelb markiert\t")
    for feld, ende in [("PAGE", " / "), ("NUMPAGES", "")]:
        element = OxmlElement("w:fldSimple"); element.set(qn("w:instr"), feld); fuss._p.append(element); fuss.add_run(ende)
    for lauf in fuss.runs: lauf.font.name = "Arial"; lauf.font.size = Pt(8)
    startseite(dokument, kurz)
    zeilen = (ORDNER / f"{stamm}.md").read_text().splitlines()
    i = next(n for n,z in enumerate(zeilen) if z.startswith("## 1 "))
    quellen = False
    kapitel = ""
    while i < len(zeilen):
        z = zeilen[i].strip(); i += 1
        if not z or z.startswith("<!--"): continue
        if z.startswith("|"):
            gruppe = [z]
            while i < len(zeilen) and zeilen[i].startswith("|"):
                gruppe.append(zeilen[i]); i += 1
            tabellenzeilen(dokument, gruppe); continue
        if z.startswith("## "):
            nummeriert = bool(re.match(r"## \d+ ", z))
            if nummeriert:
                kapitel = z[3:]
            p = dokument.add_paragraph(style="Heading 1" if nummeriert else "Heading 2")
            p.paragraph_format.page_break_before = nummeriert
            if nummeriert and not p.paragraph_format.page_break_before:
                p.paragraph_format.space_before = Pt(18)
            text_einfuegen(p, z[3:]); continue
        if z.startswith("### "):
            fortsetzungstitel = None
            if kurz == "VVT" and kapitel.startswith("2 VVT 01") and z == "### Herkunft und Empfänger":
                fortsetzungstitel = "VVT 01 Datenwege und Aufbewahrung"
            elif kurz == "VVT" and kapitel.startswith("3 VVT 02") and z == "### Herkunft und Empfänger":
                fortsetzungstitel = "VVT 02 Namensfreigabe und Widerruf"
            elif kurz == "VVT" and kapitel.startswith("4 VVT 03") and z == "### Quellen und Empfänger":
                fortsetzungstitel = "VVT 03 Zugriffe und Aufbewahrung"
            if fortsetzungstitel:
                fortsetzung = dokument.add_paragraph(fortsetzungstitel, style="Heading 1")
                fortsetzung.paragraph_format.page_break_before = True
            quellen = z[4:] in {"Quellen und Belege", "Rechtsquellen und Nachweise"}
            p = dokument.add_paragraph(style="Heading 1" if quellen else "Heading 2")
            p.paragraph_format.page_break_before = quellen
            text_einfuegen(p, z[4:]); continue
        # Die vier Unterschrifts-/Prüffelder erhalten jeweils eine eigene Zeile.
        teile = z.split(" · ") if z.startswith(("**Fachlich bestätigt durch:", "**Fachlich geprüft:")) else [z]
        for teil in teile:
            stil = "Quellen" if quellen else "List Bullet" if teil.startswith("- ") else "Normal"
            p = dokument.add_paragraph(style=stil)
            text_einfuegen(p, teil[2:] if teil.startswith("- ") else teil)
            if kurz == "VVT" and kapitel.startswith("5 VVT 04"):
                p.paragraph_format.line_spacing = 1.02
                p.paragraph_format.space_after = Pt(5)
    dokument.core_properties.title = "Technische und organisatorische Maßnahmen" if kurz == "TOM" else "Verzeichnis von Verarbeitungstätigkeiten"
    dokument.core_properties.subject = "Fortbildungsportal Schwaben"
    dokument.core_properties.author = ""; dokument.core_properties.last_modified_by = ""
    dokument.save(ORDNER / f"{stamm}.docx")


if __name__ == "__main__":
    for stamm in ["TOM_Fortbildungsportal_Schwaben", "VVT_Fortbildungsportal_Schwaben"]:
        erzeugen(stamm)
