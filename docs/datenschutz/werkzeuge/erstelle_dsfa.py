"""Erzeugt die ergänzende DSFA aus der Markdown-Quelle.

Mit dem gebündelten Python ausführen; danach alle Seiten mit render_docx.py prüfen.
Die Gestaltung nutzt dieselben Tabellen und Ausfüllmarkierungen wie TOM und VVT.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from erstelle_tom_vvt import tabelle, text_einfuegen

ORDNER = Path(__file__).resolve().parents[1]
STAMM = "DSFA_Fortbildungsportal_Schwaben"


def grundgestaltung(dokument):
    seite = dokument.sections[0]
    seite.page_width = Cm(21)
    seite.page_height = Cm(29.7)
    seite.top_margin = Cm(1.8)
    seite.bottom_margin = Cm(1.8)
    seite.left_margin = Cm(2)
    seite.right_margin = Cm(2)
    seite.header_distance = Cm(.7)
    seite.footer_distance = Cm(.8)
    for rand in dokument.styles.element.xpath(".//w:pBdr"):
        rand.getparent().remove(rand)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"]:
        stil = dokument.styles[name]
        stil.font.name = "Arial"
        stil.font.size = Pt(10.5)
        stil.font.color.rgb = RGBColor(0, 0, 0)
        stil.paragraph_format.space_after = Pt(6)
        stil.paragraph_format.line_spacing = 1.07
        stil.paragraph_format.widow_control = True
        sprache = OxmlElement("w:lang")
        sprache.set(qn("w:val"), "de-DE")
        stil.element.get_or_add_rPr().append(sprache)
    for name, groesse in [("Title", 24), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 11.5)]:
        stil = dokument.styles[name]
        stil.font.size = Pt(groesse)
        stil.font.bold = True
        stil.paragraph_format.space_before = Pt(10 if name == "Heading 2" else 0)
        stil.paragraph_format.space_after = Pt(8)
        stil.paragraph_format.keep_with_next = True
    dokument.styles["Subtitle"].font.italic = False
    for name, groesse in [("Metadaten", 9), ("Quellen", 10)]:
        stil = dokument.styles.add_style(name, 1)
        stil.base_style = dokument.styles["Normal"]
        stil.font.size = Pt(groesse)
    kopf = seite.header.paragraphs[0]
    kopf.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_ALIGN_PARAGRAPH.RIGHT)
    kopf.add_run("FORTBILDUNGSPORTAL SCHWABEN\tDSFA 01 · 1.0")
    fuss = seite.footer.paragraphs[0]
    fuss.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_ALIGN_PARAGRAPH.RIGHT)
    fuss.add_run("Behördliche Bestätigung ausstehend · Ausfüllfelder gelb\t")
    for name, trenner in [("PAGE", " / "), ("NUMPAGES", "")]:
        feld = OxmlElement("w:fldSimple")
        feld.set(qn("w:instr"), name)
        fuss._p.append(feld)
        fuss.add_run(trenner)
    for absatz in [kopf, fuss]:
        for lauf in absatz.runs:
            lauf.font.name = "Arial"
            lauf.font.size = Pt(8)
            lauf.font.color.rgb = RGBColor(0, 0, 0)


def tabellenblock(dokument, zeilen):
    inhalt = [[z.strip() for z in x.strip().strip("|").split("|")] for x in zeilen]
    inhalt = [x for x in inhalt if not all(re.fullmatch(r"[-: ]+", y) for y in x)]
    if inhalt[0][0] == "Risiko und mögliche Folgen":
        breiten = [5.0, 9.0, 3.0]
        schrift = 10
    elif inhalt[0][0] == "Wer":
        breiten = [4.0, 10.1, 2.9]
        schrift = 10
    else:
        breiten = [4.6, 12.4]
        schrift = 10
    tab = tabelle(dokument, inhalt, breiten, schrift=schrift)
    if inhalt[0][0] == "Risiko und mögliche Folgen":
        for zeile in tab.rows:
            zeile.cells[-1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def erzeugen():
    dokument = Document()
    grundgestaltung(dokument)
    zeilen = (ORDNER / f"{STAMM}.md").read_text().splitlines()
    nummer = 0
    neues_blatt = False
    quellen = False
    while nummer < len(zeilen):
        zeile = zeilen[nummer].strip()
        nummer += 1
        if not zeile:
            continue
        if zeile == "<!-- seite -->":
            neues_blatt = True
            continue
        if zeile.startswith("|"):
            gruppe = [zeile]
            while nummer < len(zeilen) and zeilen[nummer].startswith("|"):
                gruppe.append(zeilen[nummer])
                nummer += 1
            tabellenblock(dokument, gruppe)
            continue
        if zeile.startswith("# "):
            dokument.add_paragraph(zeile[2:], style="Title")
            continue
        if zeile == "Fortbildungsportal Schwaben":
            dokument.add_paragraph(zeile, style="Subtitle")
            continue
        if zeile.startswith("**DSFA 01"):
            text_einfuegen(dokument.add_paragraph(style="Metadaten"), zeile)
            continue
        if zeile.startswith("## "):
            p = dokument.add_paragraph(zeile[3:], style="Heading 1")
            p.paragraph_format.page_break_before = neues_blatt
            neues_blatt = False
            quellen = zeile.startswith("## 7 ")
            continue
        if zeile.startswith("### "):
            p = dokument.add_paragraph(zeile[4:], style="Heading 2")
            continue
        stil = "Quellen" if quellen else "Normal"
        if zeile.startswith("- "):
            zeile = zeile[2:]
        text_einfuegen(dokument.add_paragraph(style=stil), zeile)
    dokument.core_properties.title = "Datenschutzfolgenabschätzung"
    dokument.core_properties.subject = "Fortbildungsportal Schwaben"
    dokument.core_properties.author = ""
    dokument.core_properties.last_modified_by = ""
    dokument.core_properties.version = "1.0"
    dokument.save(ORDNER / f"{STAMM}.docx")


if __name__ == "__main__":
    erzeugen()
